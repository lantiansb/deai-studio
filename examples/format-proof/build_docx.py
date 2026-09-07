import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(os.environ.get('FORMAT_PROOF_OUTPUT', Path(__file__).resolve().parent))
ROOT.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.8), Inches(.8)
section.left_margin, section.right_margin = Inches(.9), Inches(.9)
for name, size in [('Normal', 12), ('Title', 25), ('Subtitle', 11), ('Heading 1', 14)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), 'Microsoft YaHei')
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + attr), None)
    color = rpr.find(qn('w:color'))
    if color is not None:
        for attr in ('themeColor', 'themeTint', 'themeShade'):
            color.attrib.pop(qn('w:' + attr), None)

normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.3
normal.space_after = Pt(8)
heading = doc.styles['Heading 1'].paragraph_format
heading.space_before, heading.space_after = Pt(18), Pt(7)
heading.keep_with_next = True
doc.add_heading('网站维护服务范围', 0)
doc.add_paragraph('虚构测试场景', style='Subtitle')
doc.add_paragraph('本说明列明网站维护的服务内容与范围边界。')
doc.add_heading('包含的服务', 1)
doc.add_paragraph('内容更新。')
doc.add_paragraph('依赖检查：每月一次。')
doc.add_heading('不包含的服务', 1)
doc.add_paragraph('新功能开发不在本维护范围内。')
doc.add_heading('紧急故障', 1)
doc.add_paragraph('紧急故障优先排查，未承诺响应时间。')
doc.add_heading('商务信息', 1)
doc.add_paragraph('本说明未列明报价和签约日期。')
doc.core_properties.title = '网站维护服务范围'
doc.core_properties.subject = '虚构测试场景'
doc.core_properties.author = 'deai-studio format proof'
doc.core_properties.keywords = 'fictional, format proof'
doc.save(ROOT / 'website-maintenance-scope.docx')
print(ROOT / 'website-maintenance-scope.docx')
